"""Import primary BLWG lichen forms into the local Meijendel databases."""

from __future__ import annotations

from collections import Counter, defaultdict
from hashlib import sha256
from pathlib import Path
import argparse
import subprocess

from openpyxl import load_workbook


ALIASES = {
    "Bacidia caligans": "Bacidina caligans",
    "Leptogium gelatinosum": "Scytinium gelatinosum",
    "Leptogium pulvinatum": "Scytinium pulvinatum",
}
SOURCE_DATASET = "Meetnet Korstmossen op de grond"
IMPORT_VERSION = "blwg-korstmos-bron-v1"
FIXED_PLOTS = frozenset((14, 15, 16, 17))
PROTECTED_TAXON = "Usnea articulata"
SCHEMA = Path(__file__).resolve().parents[1] / "database" / "blwg_korstmossen_bron_schema.sql"


class SqlExpr(str):
    """A locally constructed SQL expression, never workbook input."""


def mysql(sql: str) -> list[list[str]]:
    process = subprocess.run(
        ["/usr/local/mysql/bin/mysql", "--no-defaults", "--login-path=meijendel_root",
         "--default-character-set=utf8mb4", "--batch", "--skip-column-names"],
        input=sql, capture_output=True, text=True,
    )
    if process.returncode:
        raise RuntimeError(process.stderr.strip())
    return [line.split("\t") for line in process.stdout.splitlines()]


def sql_value(value) -> str:
    """Encode source text as UTF-8 hex, independent of MySQL quote modes."""
    if value is None:
        return "NULL"
    if isinstance(value, SqlExpr):
        return str(value)
    if isinstance(value, (int, float)):
        return str(value)
    return f"CONVERT(0x{str(value).encode('utf-8').hex()} USING utf8mb4)"


def insert_sql(table: str, columns: tuple[str, ...], rows: list[tuple]) -> str:
    if not rows:
        raise ValueError(f"no rows for {table}")
    values = ",\n".join("(" + ",".join(sql_value(value) for value in row) + ")"
                         for row in rows)
    return f"INSERT INTO {table} ({','.join(columns)}) VALUES\n{values};\n"


def workbook_rows(sheet):
    records = iter(sheet.values)
    next(records)
    return list(records)


def read_source(path: Path) -> tuple[dict, dict, list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if set(workbook.sheetnames) != {"Plotgegevens", "Bezoekgegevens", "Waarnemingen"}:
        raise ValueError("unexpected BLWG workbook tabs")
    plots = {}
    for row in workbook_rows(workbook["Plotgegevens"]):
        plot_id = int(row[0])
        if plot_id in plots:
            raise ValueError(f"duplicate plot {plot_id}")
        plots[plot_id] = dict(zip((
            "plot", "subarea", "province", "route", "x", "y", "length", "width",
            "area", "aspect", "vegetation", "management", "method", "habitat",
            "start_year", "stop_year", "wkt"), row))
    forms = {}
    for row in workbook_rows(workbook["Bezoekgegevens"]):
        kopid = int(row[5])
        if kopid in forms:
            raise ValueError(f"duplicate form {kopid}")
        forms[kopid] = {"kopid": kopid, "plot": int(row[0]),
                        "date": row[1].date().isoformat(), "report_year": row[2],
                        "observer": row[3], "note": row[4]}
    observations = []
    for excel_row, row in enumerate(workbook_rows(workbook["Waarnemingen"]), 2):
        observations.append({"row": excel_row, "plot": int(row[0]),
                             "date": row[1].date().isoformat(), "name": str(row[2]),
                             "quantity": int(row[3]), "note": row[4],
                             "kopid": int(row[5])})
    validate_source(plots, forms, observations)
    return plots, forms, observations


def map_plots(plots: dict) -> dict[int, int]:
    location_rows = mysql("SELECT meetlocatie_id,centrum_x_rd,centrum_y_rd "
                          "FROM Meijendel.korstmossen_meetlocatie "
                          "WHERE reconstructieversie='ndff-korstmos-v2';")
    if len(location_rows) != 12:
        raise ValueError("NDFF lichen location count changed")
    mapping = {}
    for plot_id, plot in plots.items():
        x, y = float(plot["x"]), float(plot["y"])
        distances = sorted(((float(row[1])-x)**2+(float(row[2])-y)**2, int(row[0]))
                           for row in location_rows)
        if distances[0][0] > 10000 or (len(distances)>1 and distances[0][0]==distances[1][0]):
            raise ValueError(f"plot {plot_id}: NDFF location not unique within 100 m")
        mapping[plot_id] = distances[0][1]
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("two BLWG plots map to one NDFF location")
    # Every original coordinate must fall within the current, established base area.
    for plot_id, plot in plots.items():
        x, y = float(plot["x"]), float(plot["y"])
        inside = mysql(
            "SELECT ST_Intersects(b.gebied_geometrie,"
            f"ST_SRID(POINT({x},{y}),28992)) "
            "FROM Meijendel.meijendel_basisgebied b "
            "JOIN Meijendel.meijendel_basisgebied_versie v "
            "ON v.basisgebiedversie_id=b.basisgebiedversie_id "
            "WHERE v.status='vastgesteld' ORDER BY v.basisgebiedversie_id DESC LIMIT 1;"
        )
        if inside != [["1"]]:
            raise ValueError(f"BLWG plot {plot_id} lies outside the current base area")
    return mapping


def ndff_source_matches(plots: dict, observations: list[dict], mapping: dict,
                        source_hash: str) -> tuple[list[dict], list[dict], list[dict], list[dict]]:
    inverse = {location: plot for plot, location in mapping.items()}
    public = mysql(
        "SELECT s.meetlocatie_id,DATE(o.periode_start),o.wetenschappelijke_naam,"
        "o.waarneming_id FROM Meijendel.korstmossen_recordselectie s "
        "JOIN Meijendel.ndff_open_waarneming o ON o.waarneming_id=s.waarneming_id "
        "WHERE s.reconstructieversie='ndff-korstmos-v2';"
    )
    public_ndff = [{"id": int(r[3]), "key": (inverse[int(r[0])],r[1],r[2])}
                   for r in public]
    public_source = [{**r,"key":(r["plot"],r["date"],ALIASES.get(r["name"],r["name"]))}
                     for r in observations if r["name"] != PROTECTED_TAXON]
    public_matches, public_links = match_results(public_source, public_ndff,
                                                  source_hash, "public")
    secure = mysql(
        "SELECT r.waarneming_id,r.periode_start,r.bron_centrum_x_rd,"
        "r.bron_centrum_y_rd,s.wetenschappelijke_naam "
        "FROM Meijendel_ndff_secure.ndff_waarneming_register r "
        "JOIN Meijendel_ndff_secure.ndff_soorten s ON s.ndff_soort_id=r.ndff_soort_id "
        "WHERE r.protocol LIKE '02.202%';"
    )
    secure_ndff = []
    for ndff_id, date, x, y, name in secure:
        if name != PROTECTED_TAXON:
            raise ValueError("unexpected protected taxon")
        distances = sorted(((float(plot["x"])-float(x))**2+
                            (float(plot["y"])-float(y))**2, plot_id)
                           for plot_id, plot in plots.items())
        if distances[0][0] > 10000:
            raise ValueError("protected NDFF observation lacks a BLWG plot")
        secure_ndff.append({"id": int(ndff_id),
                            "key": (distances[0][1], date, name)})
    secure_source = [{**r,"key":(r["plot"],r["date"],r["name"])}
                     for r in observations if r["name"] == PROTECTED_TAXON]
    secure_matches, secure_links = match_results(secure_source, secure_ndff,
                                                  source_hash, "secure")
    return public_matches, public_links, secure_matches, secure_links


def central_taxa(observations: list[dict]) -> dict[str, int]:
    public = mysql("SELECT bron_wetenschappelijke_naam,taxon_id "
                   "FROM Meijendel.taxa_bronkoppeling "
                   "WHERE bron_dataset='korstmossen_bezoek_taxon' "
                   "AND bron_versie='ndff-korstmos-v2' AND taxon_id IS NOT NULL;")
    mapping = {name: int(taxon_id) for name, taxon_id in public}
    sensitive = mysql("SELECT DISTINCT taxon_id FROM Meijendel.taxa_bronkoppeling "
                      "WHERE bron_dataset='ndff_soorten' "
                      "AND bron_wetenschappelijke_naam='Usnea articulata' "
                      "AND taxon_id IS NOT NULL;")
    if len(sensitive) != 1:
        raise ValueError("protected taxon is not uniquely linked centrally")
    mapping[PROTECTED_TAXON] = int(sensitive[0][0])
    result = {}
    for name in {r["name"] for r in observations}:
        canonical = ALIASES.get(name, name)
        if canonical not in mapping:
            raise ValueError(f"source taxon not linked centrally: {name}")
        result[name] = mapping[canonical]
    ids = ",".join(str(value) for value in sorted(set(result.values())))
    group_rows = mysql(f"SELECT t.taxon_id,t.weergavenaam,g.groep_code "
                       "FROM Meijendel.taxa t LEFT JOIN Meijendel.taxon_groepen g "
                       f"ON g.groep_id=t.groep_id WHERE t.taxon_id IN ({ids});")
    if len(group_rows) != len(set(result.values())) or any(
            name in ("NULL", "") or group != "korstmossen"
            for _, name, group in group_rows):
        raise ValueError("one or more BLWG taxa lack the central lichen group")
    return result


def taxon_link(name: str, source_hash: str) -> SqlExpr:
    return SqlExpr(
        "(SELECT koppeling_id FROM Meijendel.taxa_bronkoppeling "
        f"WHERE bron_systeem={sql_value('BLWG')} "
        f"AND bron_dataset={sql_value(SOURCE_DATASET)} "
        f"AND bron_versie={sql_value(source_hash)} "
        f"AND bron_taxon_id={sql_value('bronlabel:'+name)})"
    )


def db_match_status(status: str) -> str:
    return {"direct": "direct", "group": "groep"}[status]


def build_import_sql(path: Path, source_hash: str, plots: dict, forms: dict,
                     observations: list[dict], mapping: dict, taxa: dict,
                     public: list[dict], public_links: list[dict],
                     secure: list[dict], secure_links: list[dict]) -> tuple[str, dict]:
    if (len(plots), len(forms), len(observations), len(public), len(secure)) != (
            12, 44, 384, 364, 20):
        raise ValueError("the BLWG v2 source population changed")
    if sum(r["match_status"] == "direct" for r in public) != 210:
        raise ValueError("public one-to-one matching count changed")
    if sum(r["match_status"] == "direct" for r in secure) != 14:
        raise ValueError("protected one-to-one matching count changed")
    fixed_names = tuple(sorted({r["name"] for r in public
                                if r["plot"] in FIXED_PLOTS}))
    if len(fixed_names) != 24 or len(taxa) != 31:
        raise ValueError("BLWG source taxon scope changed")
    public_matrix = build_form_matrix(forms, public, fixed_names)
    secure_matrix = build_form_matrix(forms, secure, (PROTECTED_TAXON,))
    if (len(public_matrix),len(secure_matrix)) != (790,28):
        raise ValueError("BLWG form matrix size changed")
    if sum(r["status"]=="echte_nul" for r in public_matrix) != 426:
        raise ValueError("public source zero count changed")
    if sum(r["status"]=="echte_nul" for r in secure_matrix) != 8:
        raise ValueError("protected source zero count changed")

    parts = ["START TRANSACTION;\n"]
    parts.append(insert_sql("Meijendel.korstmossen_bronbestand", (
        "bronbestand_sha256","bronorganisatie","bestandsnaam","bronpad",
        "leverdatum","plotaantal","formulieraantal","resultaat_aantal",
        "beschermd_resultaat_aantal","importversie"), [(
        source_hash,"BLWG",path.name,str(path),"2026-10-23",12,44,384,20,
        IMPORT_VERSION)]))

    taxon_rows = []
    for name, taxon_id in sorted(taxa.items()):
        taxon_rows.append((
            "BLWG", SOURCE_DATASET, source_hash, "bronlabel:"+name, "afgeleid",
            name, "Korstmossen", source_hash, taxon_id, "kandidaat", "onbekend",
            "BLWG-bronnaam vergeleken met NDFF 02.202 en bestaande taxonkoppeling",
            IMPORT_VERSION,
            "Voorlopige centrale taxonkoppeling; BLWG-bronnaam blijft intact.",
            "Codex bronvergelijking", SqlExpr("CURRENT_TIMESTAMP(6)"),
        ))
    parts.append(insert_sql("Meijendel.taxa_bronkoppeling", (
        "bron_systeem","bron_dataset","bron_versie","bron_taxon_id",
        "bron_sleuteltype","bron_wetenschappelijke_naam","bron_soortgroep",
        "bronbestand_sha256","taxon_id","koppelstatus","taxonrelatie",
        "koppelmethode","regelversie","onderbouwing","beoordeeld_door",
        "beoordeeld_op"),taxon_rows))

    plot_rows = []
    exact_rows = []
    for plot_id, plot in sorted(plots.items()):
        sensitive_plot = plot_id in FIXED_PLOTS
        plot_rows.append((source_hash,plot_id,mapping[plot_id],"ndff-korstmos-v2",
                          plot["subarea"],plot["province"],plot["route"],
                          None if sensitive_plot else plot["x"],
                          None if sensitive_plot else plot["y"],
                          plot["length"],plot["width"],plot["area"],
                          plot["aspect"],plot["vegetation"],plot["management"],
                          plot["method"],plot["habitat"],plot["start_year"],
                          plot["stop_year"],
                          "exacte_bronlocatie_beveiligd" if sensitive_plot else "openbaar"))
        if sensitive_plot:
            exact_rows.append((source_hash,plot_id,plot["x"],plot["y"],plot["wkt"]))
    parts.append(insert_sql("Meijendel.korstmossen_bron_plot", (
        "bronbestand_sha256","bron_plot_id","ndff_meetlocatie_id",
        "ndff_reconstructieversie","subgebied","provincie","routenaam",
        "bron_x_rd","bron_y_rd","lengte_m","breedte_m","grootte_m2",
        "aspect","vegetatie","beheer","methode","begroeiingstype",
        "startjaar","stopjaar","locatie_afscherming"),plot_rows))
    parts.append(insert_sql("Meijendel.korstmossen_bron_formulier", (
        "bronbestand_sha256","kopid","bron_plot_id","velddatum",
        "rapportagejaar","waarnemer","opmerkingen","formulierstatus"), [
        (source_hash,f["kopid"],f["plot"],f["date"],f["report_year"],
         f["observer"],f["note"],
         "vast_plot_volledige_lijst" if f["plot"] in FIXED_PLOTS
         else "eenmalig_aanwezigheid")
        for f in sorted(forms.values(),key=lambda row:row["kopid"])]))

    result_cols = ("bronbestand_sha256","bronregelnummer","kopid","bronnaam",
                   "hoeveelheidscode","opmerkingen","taxon_bronkoppeling_id",
                   "ndff_koppelstatus","ndff_waarneming_id","ndff_matchgroep_sha256")
    for target, records in (("Meijendel.korstmossen_bron_resultaat",public),
                            ("Meijendel_ndff_secure.blwg_korstmossen_resultaat",secure)):
        parts.append(insert_sql(target,result_cols,[
            (source_hash,r["row"],r["kopid"],r["name"],r["quantity"],
             r["note"],taxon_link(r["name"],source_hash),db_match_status(r["match_status"]),
             r["direct_id"],r["group_hash"])
            for r in sorted(records,key=lambda item:item["row"])]))

    link_cols = ("bronbestand_sha256","ndff_matchgroep_sha256","ndff_waarneming_id")
    for target, links in (("Meijendel.korstmossen_bron_ndff_groeprecord",public_links),
                          ("Meijendel_ndff_secure.blwg_korstmossen_ndff_groeprecord",secure_links)):
        parts.append(insert_sql(target,link_cols,[
            (source_hash,r["group_hash"],r["id"]) for r in links]))

    matrix_cols = ("bronbestand_sha256","kopid","taxon_bronkoppeling_id",
                   "waarnemingsstatus","hoeveelheidscode","bronregelnummer")
    for target, rows in (("Meijendel.korstmossen_bron_formulier_taxon",public_matrix),
                         ("Meijendel_ndff_secure.blwg_korstmossen_formulier_taxon",secure_matrix)):
        parts.append(insert_sql(target,matrix_cols,[
            (source_hash,r["kopid"],taxon_link(r["name"],source_hash),
             r["status"],r["quantity"],r["row"]) for r in rows]))
    parts.append(insert_sql("Meijendel_ndff_secure.blwg_korstmossen_plot_exact", (
        "bronbestand_sha256","bron_plot_id","bron_x_rd","bron_y_rd","bron_wkt"),exact_rows))
    parts.append("COMMIT;\n")
    counts = {"plots":len(plots),"forms":len(forms),"public_results":len(public),
              "secure_results":len(secure),"public_direct":210,"secure_direct":14,
              "public_form_taxa":len(public_matrix),"public_zeros":426,
              "secure_form_taxa":len(secure_matrix),"secure_zeros":8}
    return "".join(parts), counts


def audit_import(source_hash: str, expected: dict) -> None:
    tables = {
        "plots":"Meijendel.korstmossen_bron_plot",
        "forms":"Meijendel.korstmossen_bron_formulier",
        "public_results":"Meijendel.korstmossen_bron_resultaat",
        "secure_results":"Meijendel_ndff_secure.blwg_korstmossen_resultaat",
        "public_form_taxa":"Meijendel.korstmossen_bron_formulier_taxon",
        "secure_form_taxa":"Meijendel_ndff_secure.blwg_korstmossen_formulier_taxon",
    }
    for label, table in tables.items():
        count = int(mysql(f"SELECT COUNT(*) FROM {table} "
                          f"WHERE bronbestand_sha256={sql_value(source_hash)};")[0][0])
        if count != expected[label]:
            raise ValueError(f"import audit {label}: expected {expected[label]}, got {count}")
    for label, table in (("public_zeros",tables["public_form_taxa"]),
                         ("secure_zeros",tables["secure_form_taxa"])):
        count = int(mysql(f"SELECT COUNT(*) FROM {table} "
                          f"WHERE bronbestand_sha256={sql_value(source_hash)} "
                          "AND waarnemingsstatus='echte_nul';")[0][0])
        if count != expected[label]:
            raise ValueError(f"import audit {label}: expected {expected[label]}, got {count}")
    for table, expected_count in (("Meijendel.korstmossen_bron_ndff_groeprecord",364),
                                  ("Meijendel_ndff_secure.blwg_korstmossen_ndff_groeprecord",20)):
        count = int(mysql(f"SELECT COUNT(*) FROM {table} "
                          f"WHERE bronbestand_sha256={sql_value(source_hash)};")[0][0])
        if count != expected_count:
            raise ValueError(f"NDFF group links incomplete in {table}")
    # Every imported result and zero must still reach the central taxon group.
    for table in (tables["public_form_taxa"], tables["secure_form_taxa"]):
        invalid = int(mysql(
            f"SELECT COUNT(*) FROM {table} m "
            "LEFT JOIN Meijendel.taxa_bronkoppeling k ON k.koppeling_id=m.taxon_bronkoppeling_id "
            "LEFT JOIN Meijendel.taxa t ON t.taxon_id=k.taxon_id "
            "LEFT JOIN Meijendel.taxon_groepen g ON g.groep_id=t.groep_id "
            f"WHERE m.bronbestand_sha256={sql_value(source_hash)} "
            "AND (k.koppeling_id IS NULL OR t.taxon_id IS NULL "
            "OR g.groep_code<>'korstmossen' OR g.groep_code IS NULL);"
        )[0][0])
        if invalid:
            raise ValueError(f"central taxon/group join failed in {table}: {invalid}")


def refresh_central_guards() -> None:
    """Register the new source routes in the existing whole-database taxon gate."""
    from import_external_ecology_sources import (
        CentralQueryDatabase, central_query_routes, central_query_schema_contract,
        central_query_triggers_sql,
    )
    db = CentralQueryDatabase()
    schema = db.schema()
    central_query_schema_contract(schema)
    guard_sql = central_query_triggers_sql({"routes": central_query_routes(schema)})
    mysql("USE Meijendel;\n" + guard_sql)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source",type=Path)
    parser.add_argument("--apply",action="store_true",help="insert into the local live MySQL databases")
    parser.add_argument("--refresh-taxonguards",action="store_true",
                        help="renew the audited central taxon write guards after this import")
    args = parser.parse_args()
    path = args.source.resolve(strict=True)
    source_hash = sha256(path.read_bytes()).hexdigest()
    plots,forms,observations = read_source(path)
    mapping = map_plots(plots)
    taxa = central_taxa(observations)
    public,public_links,secure,secure_links = ndff_source_matches(
        plots,observations,mapping,source_hash)
    sql,counts = build_import_sql(path,source_hash,plots,forms,observations,
                                  mapping,taxa,public,public_links,secure,secure_links)
    if args.apply:
        existing = mysql("SELECT COUNT(*) FROM information_schema.tables "
                         "WHERE table_schema='Meijendel' "
                         "AND table_name='korstmossen_bronbestand';")
        if existing == [["1"]]:
            already = mysql("SELECT bronbestand_sha256 FROM Meijendel.korstmossen_bronbestand;")
            if already:
                if already == [[source_hash]]:
                    audit_import(source_hash,counts)
                    if args.refresh_taxonguards:
                        refresh_central_guards()
                    print(f"AL_BESTAAND: BLWG-bron {source_hash}; aantallen opnieuw gecontroleerd")
                    return
                raise ValueError("different BLWG source version already imported; review before changing")
        mysql(SCHEMA.read_text(encoding="utf-8"))
        mysql(sql)
        audit_import(source_hash,counts)
        refresh_central_guards()
        print(f"GEIMPORTEERD: BLWG-bron {source_hash}")
    else:
        print(f"VOORCONTROLE: BLWG-bron {source_hash}")
    print(" ".join(f"{key}={value}" for key,value in counts.items()))


def validate_source(plots: dict, forms: dict, observations: list[dict]) -> None:
    """Reject a broken source hierarchy before any database write."""
    for form in forms.values():
        if form["plot"] not in plots:
            raise ValueError(f"form {form['kopid']}: unknown plot")
    for observation in observations:
        form = forms.get(observation["kopid"])
        if (form is None or observation["plot"] != form["plot"]
                or observation["date"] != form["date"]):
            raise ValueError(f"row {observation['row']}: form/date mismatch")


def match_results(source: list[dict], ndff: list[dict], source_hash: str,
                  domain: str) -> tuple[list[dict], list[dict]]:
    """Link unique source keys directly; preserve ambiguous matches as sets."""
    source_counts = Counter(item["key"] for item in source)
    ndff_counts = Counter(item["key"] for item in ndff)
    if source_counts != ndff_counts:
        raise ValueError("NDFF multiset mismatch")
    ndff_by_key: dict[tuple, list[int]] = defaultdict(list)
    for item in ndff:
        ndff_by_key[item["key"]].append(int(item["id"]))
    matched = []
    links = []
    seen = set()
    for item in source:
        key = item["key"]
        group_hash = sha256(repr((source_hash, domain, key)).encode()).hexdigest()
        ids = ndff_by_key[key]
        matched.append({**item, "group_hash": group_hash,
                        "match_status": "direct" if len(ids) == 1 else "group",
                        "direct_id": ids[0] if len(ids) == 1 else None})
        if key not in seen:
            links.extend({"group_hash": group_hash, "id": ndff_id} for ndff_id in ids)
            seen.add(key)
    return matched, links


def build_form_matrix(forms: dict[int, dict], observations: list[dict],
                      fixed_taxa: tuple[str, ...]) -> list[dict]:
    """Make zeros only on original complete fixed-plot forms."""
    by_form: dict[int, dict[str, dict]] = defaultdict(dict)
    for item in observations:
        if item["name"] in by_form[item["kopid"]]:
            raise ValueError(f"duplicate taxon on source form {item['kopid']}")
        by_form[item["kopid"]][item["name"]] = item
    matrix = []
    for kopid, form in sorted(forms.items()):
        found = by_form[kopid]
        names = sorted(fixed_taxa) if form["plot"] in (14, 15, 16, 17) else sorted(found)
        for name in names:
            item = found.get(name)
            matrix.append({"kopid": kopid, "name": name,
                           "status": "waargenomen" if item else "echte_nul",
                           "quantity": item["quantity"] if item else 0,
                           "row": item["row"] if item else None})
    return matrix


if __name__ == "__main__":
    main()
