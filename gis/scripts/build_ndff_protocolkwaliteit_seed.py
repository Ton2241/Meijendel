#!/usr/bin/env python3
"""Genereer de reproduceerbare protocolseed uit de beoordeelde Excel-matrix."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).parents[2]
SOURCE = ROOT / "Natuurprotocollen" / "Natuurprotocollen_gebruiksmatrix.xlsx"
OUTPUT = ROOT / "gis" / "database" / "ndff_protocolkwaliteit_seed.csv"
HEADERS = (
    "protocol_sleutel", "protocol_code", "protocol_naam", "levering_scope",
    "hoofdtype", "aanvullend_gebruik", "aanvullende_typen",
    "wetenschappelijk_gebruik", "passende_analyse",
    "benodigde_onderzoekscontext", "begrenzing", "bron_nummers", "bron_urls",
)


def text(value: object | None) -> str:
    return "" if value is None else str(value).strip()


def main() -> int:
    sheet = load_workbook(SOURCE, read_only=True, data_only=True)["Protocollen"]
    rows: list[dict[str, str]] = []
    for values in sheet.iter_rows(min_row=6, max_row=59, min_col=1, max_col=11, values_only=True):
        code, name, delivery, main_type, extra, use, analysis, context, limit, sources, urls = values
        code_value = text(code)
        key = "LOS" if code_value.casefold() == "geen code" else code_value
        extra_value = text(extra)
        extra_types = sorted(set(re.findall(r"(?<![A-Z])(TV|TA|TK|I|V)(?![A-Z])", extra_value)))
        rows.append({
            "protocol_sleutel": key,
            "protocol_code": "" if key == "LOS" else code_value,
            "protocol_naam": text(name),
            "levering_scope": "beide" if text(delivery) == "Beide leveringen" else "alleen_openbaar",
            "hoofdtype": text(main_type),
            "aanvullend_gebruik": extra_value,
            "aanvullende_typen": ",".join(extra_types),
            "wetenschappelijk_gebruik": text(use),
            "passende_analyse": text(analysis),
            "benodigde_onderzoekscontext": text(context),
            "begrenzing": text(limit),
            "bron_nummers": text(sources),
            "bron_urls": json.dumps([url for url in text(urls).splitlines() if url], ensure_ascii=False),
        })
    if len(rows) != 54:
        raise ValueError(f"Verwacht 54 protocollen, gevonden {len(rows)}")
    with OUTPUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print(f"OK: {OUTPUT} ({len(rows)} protocollen)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
