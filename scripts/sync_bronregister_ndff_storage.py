#!/usr/bin/env python3
"""Synchroniseer de fysieke NDFF-mosopslagnotitie in de bestaande DOCX."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "MEIJENDEL_BRONREGISTER.docx"
TEXT = (
    "Opslagwijziging van 23 oktober 2026. De 32.353 openbare NDFF-"
    "moswaarnemingen uit 1940–2025 en 28.643 openbare NDFF-"
    "korstmoswaarnemingen uit 1936–2025 staan nu als volledige bronregels "
    "in Meijendel.mossen_waarneming en Meijendel.korstmossen_waarneming. "
    "De overige 749.834 openbare regels blijven in ndff_open_waarneming. "
    "Een smal register verbindt de 810.830 unieke openbare identiteiten "
    "met precies één fysieke tabel; het bewaart geen tweede volledige "
    "waarneming. De vroegere ID-only-tabellen ndff_mossen en "
    "ndff_korstmossen zijn verwijderd. De afzonderlijke BLWG-formulieren "
    "en hun resultaten, de bestaande PQ-kwalificaties en de afgeschermde "
    "exacte geometrieën zijn behouden. Dit is een opslagverbetering, "
    "geen nieuwe bron of bewijs voor extra trends."
)


def main() -> None:
    doc = Document(DOCX)
    if any(paragraph.text.startswith("Opslagwijziging van 23 oktober 2026.")
           for paragraph in doc.paragraphs):
        raise RuntimeError("Opslagwijziging staat al in het DOCX; niets overschreven")
    stand = next((p for p in doc.paragraphs if p.text.strip() == "Stand: 23 oktober 2026."), None)
    if stand is None:
        raise RuntimeError("Standdatum van het bestaande bronregister niet gevonden")
    element = OxmlElement("w:p")
    stand._p.addnext(element)
    paragraph = Paragraph(element, stand._parent)
    paragraph.style = doc.styles["Normal"]
    paragraph.add_run(TEXT)
    doc.save(DOCX)
    check = Document(DOCX)
    if sum(p.text == TEXT for p in check.paragraphs) != 1:
        raise RuntimeError("DOCX-opslagnotitie niet exact eenmaal teruggelezen")
    print(f"Bijgewerkt: {DOCX}")


if __name__ == "__main__":
    main()
