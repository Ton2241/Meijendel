#!/usr/bin/env python3
"""Neem de BLWG-analysecontrole uit het canonieke Markdown-register over."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
MARKDOWN = ROOT / "docs" / "MEIJENDEL_BRONREGISTER.md"
DOCX = ROOT / "docs" / "MEIJENDEL_BRONREGISTER.docx"
PREFIX = "**Analysecontrole, 24 oktober 2026.** "
ANCHOR = "Wat kan nu? De vier vaste korstmossenplots uit 2000–2025"


def main() -> None:
    source = MARKDOWN.read_text(encoding="utf-8")
    matches = [line for line in source.splitlines() if line.startswith(PREFIX)]
    if len(matches) != 1:
        raise RuntimeError("Verwacht precies een BLWG-analysecontrole in Markdown")
    text = matches[0].replace("**", "").replace("`", "")
    document = Document(DOCX)
    stand = [p for p in document.paragraphs if p.text.startswith("Stand: ")]
    anchors = [p for p in document.paragraphs if p.text.startswith(ANCHOR)]
    if len(stand) != 1 or len(anchors) != 1:
        raise RuntimeError("Standdatum of BLWG-anker in DOCX ontbreekt of is dubbel")
    stand[0].text = "Stand: 24 oktober 2026."
    primary = next(p for p in document.paragraphs
                   if p.text.startswith("Primaire BLWG-bron, aangevuld 23 oktober"))
    old_year = ("Gebruik voor analyses de velddatum; het rapportagejaar ligt "
                "bij de toegevoegde formulieren uit 2015 en 2018 telkens één jaar eerder.")
    new_year = ("Gebruik voor analyses de velddatum: bij 40 van de 44 formulieren "
                "ligt het rapportagejaar één jaar eerder; bij de vier formulieren "
                "uit 2006 is het gelijk aan het veldjaar.")
    if old_year in primary.text:
        run = next((r for r in primary.runs if old_year in r.text), None)
        if run is None:
            raise RuntimeError("Jaarcorrectie is onverwacht over Word-runs verdeeld")
        run.text = run.text.replace(old_year, new_year)
    elif new_year not in primary.text:
        raise RuntimeError("Onbekende jaartekst in bestaande BLWG-paragraaf")

    storage = next(p for p in document.paragraphs
                   if p.text.startswith("Opslag en vervolg. De tien eerder"))
    old_storage = ("ndff_korstmossen en ndff_mossen blijven NDFF-indexen "
                   "voor alle bronhouders, niet uitsluitend BLWG.")
    new_storage = (
        "De voormalige ID-only-tabellen ndff_korstmossen en ndff_mossen "
        "zijn op 23 oktober 2026 verwijderd na fysieke verhuizing van de "
        "volledige openbare NDFF-regels naar korstmossen_waarneming en "
        "mossen_waarneming."
    )
    if old_storage in storage.text:
        start = next(i for i, run in enumerate(storage.runs)
                     if run.text == "ndff_korstmossen")
        if [r.text for r in storage.runs[start:start + 6]] != [
                "ndff_korstmossen", " ", "en", " ", "ndff_mossen", " "]:
            raise RuntimeError("Word-runs van oude indextekst zijn gewijzigd")
        tail = storage.runs[start + 6]
        old_tail = ("blijven NDFF-indexen voor alle bronhouders, "
                    "niet uitsluitend BLWG. ")
        if not tail.text.startswith(old_tail):
            raise RuntimeError("Oude indextekst heeft onbekend vervolg")
        storage.runs[start].text = new_storage + " "
        for run in storage.runs[start + 1:start + 6]:
            run.text = ""
        tail.text = tail.text[len(old_tail):]
    elif new_storage not in storage.text:
        raise RuntimeError("Onbekende indextekst in bestaande BLWG-paragraaf")
    existing = [p for p in document.paragraphs
                if p.text.startswith("Analysecontrole, 24 oktober 2026.")]
    if len(existing) > 1:
        raise RuntimeError("BLWG-analysecontrole dubbel in DOCX")
    if existing:
        existing[0].text = text
    else:
        element = OxmlElement("w:p")
        anchors[0]._p.addnext(element)
        paragraph = Paragraph(element, anchors[0]._parent)
        paragraph.style = document.styles["Normal"]
        paragraph.add_run(text)
    document.save(DOCX)
    check = Document(DOCX)
    if sum(p.text == text for p in check.paragraphs) != 1:
        raise RuntimeError("DOCX bevat de Markdown-tekst niet exact eenmaal")
    if new_year not in "\n".join(p.text for p in check.paragraphs):
        raise RuntimeError("Jaarcorrectie ontbreekt in DOCX")
    if new_storage not in "\n".join(p.text for p in check.paragraphs):
        raise RuntimeError("Indexcorrectie ontbreekt in DOCX")
    print(f"Gesynchroniseerd: {DOCX}")


if __name__ == "__main__":
    main()
